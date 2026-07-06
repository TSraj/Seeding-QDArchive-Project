"""
Step 5: Extract XLSX Table

Queries the classification database and exports a spreadsheet with the
required columns:
    repository_id, project_type, project_title, primary_class,
    secondary_class, no_project_files

Usage:
    python -m src.classification.step5_extract_table
"""

import sqlite3
import os

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("ERROR: openpyxl is required. Install with:  pip install openpyxl")
    raise


def run(db_path: str, output_path: str):
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()

    # Query: join projects → project_classifications + file count
    # Left join so that projects without a classification still appear.
    query = """
    SELECT
        p.repository_id,
        p.type            AS project_type,
        p.title           AS project_title,
        pc.isic_division || ' - ' || pc.isic_division_name AS primary_class,
        CASE
            WHEN pc.secondary_isic_division IS NOT NULL AND pc.secondary_isic_division != ''
            THEN pc.secondary_isic_division || ' - ' || pc.secondary_isic_division_name
            ELSE ''
        END AS secondary_class,
        COALESCE(fc.file_count, 0) AS no_project_files
    FROM projects p
    LEFT JOIN project_classifications pc ON pc.project_id = p.id
    LEFT JOIN (
        SELECT project_id, COUNT(*) AS file_count
        FROM files
        GROUP BY project_id
    ) fc ON fc.project_id = p.id
    ORDER BY p.repository_id, p.type, p.title
    """

    cursor.execute(query)
    rows = cursor.fetchall()
    conn.close()

    # --- Build the workbook ---
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Project Classifications"

    # Column headers
    headers = [
        "repository_id",
        "project_type",
        "project_title",
        "primary_class",
        "secondary_class",
        "no_project_files",
    ]

    # Styling
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # Write headers
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    # Write data rows
    alt_fill = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
    data_font = Font(name="Calibri", size=10)

    for row_idx, row in enumerate(rows, start=2):
        values = [
            row["repository_id"],
            row["project_type"],
            row["project_title"],
            row["primary_class"] if row["primary_class"] else "",
            row["secondary_class"] if row["secondary_class"] else "",
            row["no_project_files"],
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = data_font
            cell.border = thin_border
            if row_idx % 2 == 0:
                cell.fill = alt_fill

    # Auto-fit column widths (approximate)
    for col_idx in range(1, len(headers) + 1):
        col_letter = get_column_letter(col_idx)
        max_length = len(headers[col_idx - 1])
        for row_idx in range(2, min(len(rows) + 2, 200)):  # sample first 200 rows
            cell_value = str(ws.cell(row=row_idx, column=col_idx).value or "")
            max_length = max(max_length, len(cell_value))
        ws.column_dimensions[col_letter].width = min(max_length + 4, 60)

    # Freeze header row
    ws.freeze_panes = "A2"

    # Auto-filter
    ws.auto_filter.ref = ws.dimensions

    # Save
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)

    print(f"\n✅ Step 5: XLSX table exported successfully!")
    print(f"   → {output_path}")
    print(f"   → {len(rows)} projects written ({len(headers)} columns)")


if __name__ == "__main__":
    script_dir = os.path.dirname(os.path.abspath(__file__))
    project_root = os.path.join(script_dir, "..", "..")
    db_path = os.path.join(project_root, "23359384-sq26-classification.db")
    output_path = os.path.join(project_root, "reports", "project_classification_table.xlsx")

    run(db_path, output_path)
